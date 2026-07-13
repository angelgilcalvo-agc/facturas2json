import openpyxl

wb = openpyxl.load_workbook("data/interim/ANA_FEIJOO_CALCULO_2T_2026.xlsx")
sheet = wb["Resumen ejecutivo"]

for c in range(1, 7):
    sheet.cell(row=11, column=c, value=None)

sheet.cell(row=12, column=1, value="Indicador 2T")
sheet.cell(row=12, column=2, value="Importe/Base")
sheet.cell(row=12, column=3, value="IVA Soportado")
sheet.cell(row=12, column=4, value="IVA Devengado ISP")
sheet.cell(row=12, column=5, value="Total")
sheet.cell(row=12, column=6, value="Notas")

sheet.cell(row=13, column=2, value="='IVA por tipos'!C3 + 280.00")
sheet.cell(row=13, column=3, value="='IVA por tipos'!E3 + 58.80")
sheet.cell(row=13, column=4, value=None)
sheet.cell(row=13, column=5, value="=B13+C13")
sheet.cell(row=13, column=6, value="Incluye AUCTANE, CIFRA, Amazon, Amazon Ads España, Eurogrow, IONOS, Finetwork, Digi, Navapack, PlasticForte, Leroy Merlin, alquiler, portes 21% y facturas mixtas")

sheet.cell(row=14, column=2, value="='IVA por tipos'!C4")
sheet.cell(row=14, column=3, value=0.0)
sheet.cell(row=14, column=4, value=None)
sheet.cell(row=14, column=5, value="=B14")
sheet.cell(row=14, column=6, value="TGSS autónomos, comunidad, seguros de préstamo y otros sin IVA; revisar deducibilidad/afectación")

sheet.cell(row=15, column=2, value="='IVA por tipos'!C5")
sheet.cell(row=15, column=3, value=None)
sheet.cell(row=15, column=4, value="='IVA por tipos'!F5")
sheet.cell(row=15, column=5, value="=B15")
sheet.cell(row=15, column=6, value="ARTEVASI Portugal y VIVEIROS CULTIFLOR; ISP estimada al 21%")

sheet.cell(row=16, column=2, value="='IVA por tipos'!C6")
sheet.cell(row=16, column=3, value=None)
sheet.cell(row=16, column=4, value="='IVA por tipos'!F6")
sheet.cell(row=16, column=5, value="=B16")
sheet.cell(row=16, column=6, value="Amazon Ads Italia/Francia y Stripe Irlanda con reverse charge; ISP estimada al 21%")

sheet.cell(row=17, column=2, value="='IVA por tipos'!C7")
sheet.cell(row=17, column=3, value="='IVA por tipos'!E7")
sheet.cell(row=17, column=4, value=None)
sheet.cell(row=17, column=5, value="=B17+C17")
sheet.cell(row=17, column=6, value="Iberdrola, Siro, Projar y Ricardo Terán; revisar afectación/deducibilidad donde proceda")

sheet.cell(row=18, column=2, value="='IVA por tipos'!C9")
sheet.cell(row=18, column=3, value=None)
sheet.cell(row=18, column=4, value=None)
sheet.cell(row=18, column=5, value="=B18")
sheet.cell(row=18, column=6, value="Alibaba/Wuhan China; no se incorpora IVA importación sin DUA")

sheet.cell(row=19, column=2, value="=SUM(B13:B18)")
sheet.cell(row=19, column=3, value="=SUM(C13,C17)")
sheet.cell(row=19, column=4, value="=SUM(D15:D16)")
sheet.cell(row=19, column=5, value="=SUM(E13:E18)")
sheet.cell(row=19, column=6, value="Incluye todas las facturas y documentos del 2T; excluye importes fuera de periodo; ISP UE queda como estimación")

wb.save("data/interim/ANA_FEIJOO_CALCULO_2T_2026.xlsx")
print("Fixed Resumen ejecutivo.")
