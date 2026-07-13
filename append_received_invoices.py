import openpyxl
import datetime

# Load the working workbook
wb = openpyxl.load_workbook("data/interim/ANA_FEIJOO_CALCULO_2T_2026.xlsx")

# 20 new received invoices extracted earlier
# Map each file's precise values to put in "Facturas recibidas" and "Documentos procesados"
new_invoices = [
    # Alquiler
    {
        "Archivo": "Alquiler JUNIO 03-26 ANA BELEN FEIJOO.pdf",
        "Fecha": datetime.date(2026, 6, 4),
        "Nº factura/documento": "3",
        "Proveedor": "ROSA FERNANDEZ DA BOUZA",
        "NIF/CIF proveedor": "34934884T",
        "País": "España",
        "Tipo operación": "Alquiler de local comercial sujeto a retención",
        "Descripción": "Alquiler local comercial Camino Caneiro 11",
        "Base imponible EUR": 280.00,
        "% IVA": 0.21,
        "IVA soportado EUR": 58.80,
        "Total EUR": 285.60, # Note: Total = Base + IVA - Ret = 280 + 58.80 - 53.20 = 285.60
        "Moneda original": "EUR",
        "Importe original": 285.60,
        "Deducibilidad": 1.0,
        "% Retención": 0.19,
        "Retención EUR": 53.20,
        "Clasificación fiscal": "Alquiler de local comercial sujeto a retención",
        "Estado revisión": "Revisado",
        "Observaciones": "Alquiler de local comercial sujeto a retención. Base: 280,00 €, IVA 21%: 58,80 €, Retención IRPF 19%: 53,20 €.",
        "Documento origen": "Alquiler JUNIO 03-26 ANA BELEN FEIJOO.pdf",
        "Incluye en 2T": "Sí"
    },
    # Amazon invoices / credit notes
    {
        "Archivo": "ES-AEU-2026-484690.pdf",
        "Fecha": datetime.date(2026, 6, 30),
        "Nº factura/documento": "ES-AEU-2026-484690",
        "Proveedor": "Amazon EU S.à r.l., Sucursal en España",
        "NIF/CIF proveedor": "ESW0184081H",
        "País": "España",
        "Tipo operación": "Compra/servicio nacional con IVA",
        "Descripción": "Tarifas de logística de Amazon",
        "Base imponible EUR": 4065.81,
        "% IVA": 0.21,
        "IVA soportado EUR": 853.82,
        "Total EUR": 4919.63,
        "Moneda original": "EUR",
        "Importe original": 4919.63,
        "Deducibilidad": 1.0,
        "% Retención": 0.0,
        "Retención EUR": 0.0,
        "Clasificación fiscal": "Gasto con IVA soportado nacional",
        "Estado revisión": "Revisado",
        "Observaciones": "Tarifas de logística de Amazon con IVA 21%.",
        "Documento origen": "ES-AEU-2026-484690.pdf",
        "Incluye en 2T": "Sí"
    },
    {
        "Archivo": "ES-AEU-2026-484703.pdf",
        "Fecha": datetime.date(2026, 6, 30),
        "Nº factura/documento": "ES-AEU-2026-484703",
        "Proveedor": "Amazon EU S.à r.l., Sucursal en España",
        "NIF/CIF proveedor": "ESW0184081H",
        "País": "España",
        "Tipo operación": "Compra/servicio nacional con IVA",
        "Descripción": "Tarifas de logística de Amazon",
        "Base imponible EUR": 1222.46,
        "% IVA": 0.21,
        "IVA soportado EUR": 256.72,
        "Total EUR": 1479.18,
        "Moneda original": "EUR",
        "Importe original": 1479.18,
        "Deducibilidad": 1.0,
        "% Retención": 0.0,
        "Retención EUR": 0.0,
        "Clasificación fiscal": "Gasto con IVA soportado nacional",
        "Estado revisión": "Revisado",
        "Observaciones": "Tarifas de logística de Amazon con IVA 21%.",
        "Documento origen": "ES-AEU-2026-484703.pdf",
        "Incluye en 2T": "Sí"
    },
    {
        "Archivo": "ES-AEU-2026-497086.pdf",
        "Fecha": datetime.date(2026, 6, 30),
        "Nº factura/documento": "ES-AEU-2026-497086",
        "Proveedor": "Amazon EU S.à r.l., Sucursal en España",
        "NIF/CIF proveedor": "ESW0184081H",
        "País": "España",
        "Tipo operación": "Operación exenta/no sujeta sin IVA",
        "Descripción": "Tarifas de logística de Amazon (exentas)",
        "Base imponible EUR": 28.88,
        "% IVA": 0.0,
        "IVA soportado EUR": 0.0,
        "Total EUR": 28.88,
        "Moneda original": "EUR",
        "Importe original": 28.88,
        "Deducibilidad": 1.0,
        "% Retención": 0.0,
        "Retención EUR": 0.0,
        "Clasificación fiscal": "Operación no sujeta/exenta sin IVA",
        "Estado revisión": "Revisado",
        "Observaciones": "Tarifas de logística de Amazon exentas (Art. 69, 21.5 o 64 LIVA).",
        "Documento origen": "ES-AEU-2026-497086.pdf",
        "Incluye en 2T": "Sí"
    },
    {
        "Archivo": "ES-AEU-2026-502268.pdf",
        "Fecha": datetime.date(2026, 6, 30),
        "Nº factura/documento": "ES-AEU-2026-502268",
        "Proveedor": "Amazon EU S.à r.l., Sucursal en España",
        "NIF/CIF proveedor": "ESW0184081H",
        "País": "España",
        "Tipo operación": "Compra/servicio nacional con IVA",
        "Descripción": "Tarifas de logística de Amazon",
        "Base imponible EUR": 397.44,
        "% IVA": 0.21,
        "IVA soportado EUR": 83.46,
        "Total EUR": 480.90,
        "Moneda original": "EUR",
        "Importe original": 480.90,
        "Deducibilidad": 1.0,
        "% Retención": 0.0,
        "Retención EUR": 0.0,
        "Clasificación fiscal": "Gasto con IVA soportado nacional",
        "Estado revisión": "Revisado",
        "Observaciones": "Tarifas de logística de Amazon con IVA 21%. Factura en francés.",
        "Documento origen": "ES-AEU-2026-502268.pdf",
        "Incluye en 2T": "Sí"
    },
    {
        "Archivo": "ES-AEU-2026-502276.pdf",
        "Fecha": datetime.date(2026, 6, 30),
        "Nº factura/documento": "ES-AEU-2026-502276",
        "Proveedor": "Amazon EU S.à r.l., Sucursal en España",
        "NIF/CIF proveedor": "ESW0184081H",
        "País": "España",
        "Tipo operación": "Compra/servicio nacional con IVA",
        "Descripción": "Tarifas de logística de Amazon",
        "Base imponible EUR": 70.65,
        "% IVA": 0.21,
        "IVA soportado EUR": 14.84,
        "Total EUR": 85.49,
        "Moneda original": "EUR",
        "Importe original": 85.49,
        "Deducibilidad": 1.0,
        "% Retención": 0.0,
        "Retención EUR": 0.0,
        "Clasificación fiscal": "Gasto con IVA soportado nacional",
        "Estado revisión": "Revisado",
        "Observaciones": "Tarifas de logística de Amazon con IVA 21%. Factura en francés.",
        "Documento origen": "ES-AEU-2026-502276.pdf",
        "Incluye en 2T": "Sí"
    },
    {
        "Archivo": "ES-AEU-2026-513630.pdf",
        "Fecha": datetime.date(2026, 6, 30),
        "Nº factura/documento": "ES-AEU-2026-513630",
        "Proveedor": "Amazon EU S.à r.l., Sucursal en España",
        "NIF/CIF proveedor": "ESW0184081H",
        "País": "España",
        "Tipo operación": "Compra/servicio nacional con IVA",
        "Descripción": "Tarifas de logística de Amazon",
        "Base imponible EUR": 84.13,
        "% IVA": 0.21,
        "IVA soportado EUR": 17.67,
        "Total EUR": 101.80,
        "Moneda original": "EUR",
        "Importe original": 101.80,
        "Deducibilidad": 1.0,
        "% Retención": 0.0,
        "Retención EUR": 0.0,
        "Clasificación fiscal": "Gasto con IVA soportado nacional",
        "Estado revisión": "Revisado",
        "Observaciones": "Tarifas de logística de Amazon con IVA 21%. Factura en italiano.",
        "Documento origen": "ES-AEU-2026-513630.pdf",
        "Incluye en 2T": "Sí"
    },
    {
        "Archivo": "ES-AEU-2026-513631.pdf",
        "Fecha": datetime.date(2026, 6, 30),
        "Nº factura/documento": "ES-AEU-2026-513631",
        "Proveedor": "Amazon EU S.à r.l., Sucursal en España",
        "NIF/CIF proveedor": "ESW0184081H",
        "País": "España",
        "Tipo operación": "Compra/servicio nacional con IVA",
        "Descripción": "Tarifas de logística de Amazon",
        "Base imponible EUR": 16.84,
        "% IVA": 0.21,
        "IVA soportado EUR": 3.54,
        "Total EUR": 20.38,
        "Moneda original": "EUR",
        "Importe original": 20.38,
        "Deducibilidad": 1.0,
        "% Retención": 0.0,
        "Retención EUR": 0.0,
        "Clasificación fiscal": "Gasto con IVA soportado nacional",
        "Estado revisión": "Revisado",
        "Observaciones": "Tarifas de logística de Amazon con IVA 21%. Factura en italiano.",
        "Documento origen": "ES-AEU-2026-513631.pdf",
        "Incluye en 2T": "Sí"
    },
    {
        "Archivo": "ES-AEU-2026-549585.pdf",
        "Fecha": datetime.date(2026, 6, 30),
        "Nº factura/documento": "ES-AEU-2026-549585",
        "Proveedor": "Amazon EU S.à r.l., Sucursal en España",
        "NIF/CIF proveedor": "ESW0184081H",
        "País": "España",
        "Tipo operación": "Compra/servicio nacional con IVA",
        "Descripción": "Tarifas de logística de Amazon",
        "Base imponible EUR": 4.62,
        "% IVA": 0.21,
        "IVA soportado EUR": 0.97,
        "Total EUR": 5.59,
        "Moneda original": "EUR",
        "Importe original": 5.59,
        "Deducibilidad": 1.0,
        "% Retención": 0.0,
        "Retención EUR": 0.0,
        "Clasificación fiscal": "Gasto con IVA soportado nacional",
        "Estado revisión": "Revisado",
        "Observaciones": "Tarifas de logística de Amazon con IVA 21%. Factura en alemán.",
        "Documento origen": "ES-AEU-2026-549585.pdf",
        "Incluye en 2T": "Sí"
    },
    {
        "Archivo": "ES-AEU-2026-573889.pdf",
        "Fecha": datetime.date(2026, 6, 30),
        "Nº factura/documento": "ES-AEU-2026-573889",
        "Proveedor": "Amazon EU S.à r.l., Sucursal en España",
        "NIF/CIF proveedor": "ESW0184081H",
        "País": "España",
        "Tipo operación": "Compra/servicio nacional con IVA",
        "Descripción": "Tarifas de logística de Amazon",
        "Base imponible EUR": 1429.74,
        "% IVA": 0.21,
        "IVA soportado EUR": 300.25,
        "Total EUR": 1729.99,
        "Moneda original": "EUR",
        "Importe original": 1729.99,
        "Deducibilidad": 1.0,
        "% Retención": 0.0,
        "Retención EUR": 0.0,
        "Clasificación fiscal": "Gasto con IVA soportado nacional",
        "Estado revisión": "Revisado",
        "Observaciones": "Tarifas de logística de Amazon con IVA 21%.",
        "Documento origen": "ES-AEU-2026-573889.pdf",
        "Incluye en 2T": "Sí"
    },
    {
        "Archivo": "ES-AEU-2026-573890.pdf",
        "Fecha": datetime.date(2026, 6, 30),
        "Nº factura/documento": "ES-AEU-2026-573890",
        "Proveedor": "Amazon EU S.à r.l., Sucursal en España",
        "NIF/CIF proveedor": "ESW0184081H",
        "País": "España",
        "Tipo operación": "Compra/servicio nacional con IVA",
        "Descripción": "Tarifas de logística de Amazon",
        "Base imponible EUR": 407.98,
        "% IVA": 0.21,
        "IVA soportado EUR": 85.67,
        "Total EUR": 493.65,
        "Moneda original": "EUR",
        "Importe original": 493.65,
        "Deducibilidad": 1.0,
        "% Retención": 0.0,
        "Retención EUR": 0.0,
        "Clasificación fiscal": "Gasto con IVA soportado nacional",
        "Estado revisión": "Revisado",
        "Observaciones": "Tarifas de logística de Amazon con IVA 21%.",
        "Documento origen": "ES-AEU-2026-573890.pdf",
        "Incluye en 2T": "Sí"
    },
    {
        "Archivo": "ES-AEU-2026-574007.pdf",
        "Fecha": datetime.date(2026, 6, 30),
        "Nº factura/documento": "ES-AEU-2026-574007",
        "Proveedor": "Amazon EU S.à r.l., Sucursal en España",
        "NIF/CIF proveedor": "ESW0184081H",
        "País": "España",
        "Tipo operación": "Compra/servicio nacional con IVA",
        "Descripción": "Tarifas de logística de Amazon",
        "Base imponible EUR": 88.28,
        "% IVA": 0.21,
        "IVA soportado EUR": 18.54,
        "Total EUR": 106.82,
        "Moneda original": "EUR",
        "Importe original": 106.82,
        "Deducibilidad": 1.0,
        "% Retención": 0.0,
        "Retención EUR": 0.0,
        "Clasificación fiscal": "Gasto con IVA soportado nacional",
        "Estado revisión": "Revisado",
        "Observaciones": "Tarifas de logística de Amazon con IVA 21%. Factura en francés.",
        "Documento origen": "ES-AEU-2026-574007.pdf",
        "Incluye en 2T": "Sí"
    },
    {
        "Archivo": "ES-AEU-2026-574008.pdf",
        "Fecha": datetime.date(2026, 6, 30),
        "Nº factura/documento": "ES-AEU-2026-574008",
        "Proveedor": "Amazon EU S.à r.l., Sucursal en España",
        "NIF/CIF proveedor": "ESW0184081H",
        "País": "España",
        "Tipo operación": "Compra/servicio nacional con IVA",
        "Descripción": "Tarifas de logística de Amazon",
        "Base imponible EUR": 17.98,
        "% IVA": 0.21,
        "IVA soportado EUR": 3.77,
        "Total EUR": 21.75,
        "Moneda original": "EUR",
        "Importe original": 21.75,
        "Deducibilidad": 1.0,
        "% Retención": 0.0,
        "Retención EUR": 0.0,
        "Clasificación fiscal": "Gasto con IVA soportado nacional",
        "Estado revisión": "Revisado",
        "Observaciones": "Tarifas de logística de Amazon con IVA 21%. Factura en francés.",
        "Documento origen": "ES-AEU-2026-574008.pdf",
        "Incluye en 2T": "Sí"
    },
    {
        "Archivo": "ES-AEU-2026-574355.pdf",
        "Fecha": datetime.date(2026, 6, 30),
        "Nº factura/documento": "ES-AEU-2026-574355",
        "Proveedor": "Amazon EU S.à r.l., Sucursal en España",
        "NIF/CIF proveedor": "ESW0184081H",
        "País": "España",
        "Tipo operación": "Compra/servicio nacional con IVA",
        "Descripción": "Tarifas de logística de Amazon",
        "Base imponible EUR": 24.71,
        "% IVA": 0.21,
        "IVA soportado EUR": 5.19,
        "Total EUR": 29.90,
        "Moneda original": "EUR",
        "Importe original": 29.90,
        "Deducibilidad": 1.0,
        "% Retención": 0.0,
        "Retención EUR": 0.0,
        "Clasificación fiscal": "Gasto con IVA soportado nacional",
        "Estado revisión": "Revisado",
        "Observaciones": "Tarifas de logística de Amazon con IVA 21%. Factura en italiano.",
        "Documento origen": "ES-AEU-2026-574355.pdf",
        "Incluye en 2T": "Sí"
    },
    {
        "Archivo": "ES-AEU-2026-574356.pdf",
        "Fecha": datetime.date(2026, 6, 30),
        "Nº factura/documento": "ES-AEU-2026-574356",
        "Proveedor": "Amazon EU S.à r.l., Sucursal en España",
        "NIF/CIF proveedor": "ESW0184081H",
        "País": "España",
        "Tipo operación": "Compra/servicio nacional con IVA",
        "Descripción": "Tarifas de logística de Amazon",
        "Base imponible EUR": 7.37,
        "% IVA": 0.21,
        "IVA soportado EUR": 1.55,
        "Total EUR": 8.92,
        "Moneda original": "EUR",
        "Importe original": 8.92,
        "Deducibilidad": 1.0,
        "% Retención": 0.0,
        "Retención EUR": 0.0,
        "Clasificación fiscal": "Gasto con IVA soportado nacional",
        "Estado revisión": "Revisado",
        "Observaciones": "Tarifas de logística de Amazon con IVA 21%. Factura en italiano.",
        "Documento origen": "ES-AEU-2026-574356.pdf",
        "Incluye en 2T": "Sí"
    },
    {
        "Archivo": "ES-AEU-2026-576839.pdf",
        "Fecha": datetime.date(2026, 6, 30),
        "Nº factura/documento": "ES-AEU-2026-576839",
        "Proveedor": "Amazon EU S.à r.l., Sucursal en España",
        "NIF/CIF proveedor": "ESW0184081H",
        "País": "España",
        "Tipo operación": "Operación exenta/no sujeta sin IVA",
        "Descripción": "Tarifas de logística de Amazon (exentas)",
        "Base imponible EUR": 5.77,
        "% IVA": 0.0,
        "IVA soportado EUR": 0.0,
        "Total EUR": 5.77,
        "Moneda original": "EUR",
        "Importe original": 5.77,
        "Deducibilidad": 1.0,
        "% Retención": 0.0,
        "Retención EUR": 0.0,
        "Clasificación fiscal": "Operación no sujeta/exenta sin IVA",
        "Estado revisión": "Revisado",
        "Observaciones": "Tarifas de logística de Amazon exentas (Art. 69, 21.5 o 64 LIVA).",
        "Documento origen": "ES-AEU-2026-576839.pdf",
        "Incluye en 2T": "Sí"
    },
    # Credit notes
    {
        "Archivo": "ES-CN-AEU-2026-134400.pdf",
        "Fecha": datetime.date(2026, 6, 30),
        "Nº factura/documento": "ES-CN-AEU-2026-134400",
        "Proveedor": "Amazon EU S.à r.l., Sucursal en España",
        "NIF/CIF proveedor": "ESW0184081H",
        "País": "España",
        "Tipo operación": "Rectificación de factura nacional (negativa)",
        "Descripción": "Tarifas reembolsadas",
        "Base imponible EUR": -47.11,
        "% IVA": 0.21,
        "IVA soportado EUR": -9.89,
        "Total EUR": -57.00,
        "Moneda original": "EUR",
        "Importe original": -57.00,
        "Deducibilidad": 1.0,
        "% Retención": 0.0,
        "Retención EUR": 0.0,
        "Clasificación fiscal": "Rectificación menor gasto e IVA soportado",
        "Estado revisión": "Revisado",
        "Observaciones": "Nota de crédito de tarifas de logística de Amazon.",
        "Documento origen": "ES-CN-AEU-2026-134400.pdf",
        "Incluye en 2T": "Sí"
    },
    {
        "Archivo": "ES-CN-AEU-2026-142678.pdf",
        "Fecha": datetime.date(2026, 6, 30),
        "Nº factura/documento": "ES-CN-AEU-2026-142678",
        "Proveedor": "Amazon EU S.à r.l., Sucursal en España",
        "NIF/CIF proveedor": "ESW0184081H",
        "País": "España",
        "Tipo operación": "Rectificación de factura nacional (negativa)",
        "Descripción": "Tarifas reembolsadas",
        "Base imponible EUR": -2.42,
        "% IVA": 0.21,
        "IVA soportado EUR": -0.51,
        "Total EUR": -2.93,
        "Moneda original": "EUR",
        "Importe original": -2.93,
        "Deducibilidad": 1.0,
        "% Retención": 0.0,
        "Retención EUR": 0.0,
        "Clasificación fiscal": "Rectificación menor gasto e IVA soportado",
        "Estado revisión": "Revisado",
        "Observaciones": "Nota de crédito de tarifas de logística de Amazon.",
        "Documento origen": "ES-CN-AEU-2026-142678.pdf",
        "Incluye en 2T": "Sí"
    },
    {
        "Archivo": "ES-CN-AEU-2026-154550.pdf",
        "Fecha": datetime.date(2026, 6, 30),
        "Nº factura/documento": "ES-CN-AEU-2026-154550",
        "Proveedor": "Amazon EU S.à r.l., Sucursal en España",
        "NIF/CIF proveedor": "ESW0184081H",
        "País": "España",
        "Tipo operación": "Rectificación de factura nacional (negativa)",
        "Descripción": "Tarifas reembolsadas",
        "Base imponible EUR": -4.50,
        "% IVA": 0.21,
        "IVA soportado EUR": -0.94,
        "Total EUR": -5.44,
        "Moneda original": "EUR",
        "Importe original": -5.44,
        "Deducibilidad": 1.0,
        "% Retención": 0.0,
        "Retención EUR": 0.0,
        "Clasificación fiscal": "Rectificación menor gasto e IVA soportado",
        "Estado revisión": "Revisado",
        "Observaciones": "Nota de crédito de tarifas de logística de Amazon.",
        "Documento origen": "ES-CN-AEU-2026-154550.pdf",
        "Incluye en 2T": "Sí"
    },
    {
        "Archivo": "ES-CN-AEU-2026-155131.pdf",
        "Fecha": datetime.date(2026, 6, 30),
        "Nº factura/documento": "ES-CN-AEU-2026-155131",
        "Proveedor": "Amazon EU S.à r.l., Sucursal en España",
        "NIF/CIF proveedor": "ESW0184081H",
        "País": "España",
        "Tipo operación": "Rectificación de factura nacional (negativa)",
        "Descripción": "Tarifas reembolsadas",
        "Base imponible EUR": -3.85,
        "% IVA": 0.0,
        "IVA soportado EUR": 0.0,
        "Total EUR": -3.85,
        "Moneda original": "EUR",
        "Importe original": -3.85,
        "Deducibilidad": 1.0,
        "% Retención": 0.0,
        "Retención EUR": 0.0,
        "Clasificación fiscal": "Operación no sujeta/exenta sin IVA",
        "Estado revisión": "Revisado",
        "Observaciones": "Nota de crédito de tarifas de logística de Amazon exenta.",
        "Documento origen": "ES-CN-AEU-2026-155131.pdf",
        "Incluye en 2T": "Sí"
    }
]

# Sheet: Facturas recibidas
sheet_rec = wb["Facturas recibidas"]
max_row_rec = sheet_rec.max_row
current_row = max_row_rec + 1

while current_row > 1:
    val = sheet_rec.cell(row=current_row-1, column=1).value
    if val is not None:
        break
    current_row -= 1

for idx, inv in enumerate(new_invoices):
    sheet_rec.cell(row=current_row, column=1, value="ANA BELEN FEIJOO VIDUEIRA")
    sheet_rec.cell(row=current_row, column=2, value="34975084L")
    sheet_rec.cell(row=current_row, column=3, value="2T 2026")
    sheet_rec.cell(row=current_row, column=4, value=inv["Fecha"])
    sheet_rec.cell(row=current_row, column=5, value=inv["Nº factura/documento"])
    sheet_rec.cell(row=current_row, column=6, value=inv["Proveedor"])
    sheet_rec.cell(row=current_row, column=7, value=inv["NIF/CIF proveedor"])
    sheet_rec.cell(row=current_row, column=8, value=inv["País"])
    sheet_rec.cell(row=current_row, column=9, value=inv["Tipo operación"])
    sheet_rec.cell(row=current_row, column=10, value=inv["Descripción"])
    sheet_rec.cell(row=current_row, column=11, value=inv["Base imponible EUR"])
    sheet_rec.cell(row=current_row, column=12, value=inv["% IVA"])
    sheet_rec.cell(row=current_row, column=13, value=inv["IVA soportado EUR"])
    sheet_rec.cell(row=current_row, column=14, value=inv["Total EUR"])
    sheet_rec.cell(row=current_row, column=15, value=inv["Moneda original"])
    sheet_rec.cell(row=current_row, column=16, value=inv["Importe original"])
    sheet_rec.cell(row=current_row, column=17, value=inv["Deducibilidad"])
    sheet_rec.cell(row=current_row, column=18, value=inv["% Retención"])
    sheet_rec.cell(row=current_row, column=19, value=inv["Retención EUR"])
    sheet_rec.cell(row=current_row, column=20, value=inv["Clasificación fiscal"])
    sheet_rec.cell(row=current_row, column=21, value=inv["Estado revisión"])
    sheet_rec.cell(row=current_row, column=22, value=inv["Observaciones"])
    sheet_rec.cell(row=current_row, column=23, value=inv["Documento origen"])
    sheet_rec.cell(row=current_row, column=24, value=inv["Incluye en 2T"])
    current_row += 1

# Sheet: Documentos procesados
sheet_proc = wb["Documentos procesados"]
max_row_proc = sheet_proc.max_row
current_row_proc = max_row_proc + 1
while current_row_proc > 1:
    val = sheet_proc.cell(row=current_row_proc-1, column=1).value
    if val is not None:
        break
    current_row_proc -= 1

for idx, inv in enumerate(new_invoices):
    doc_num = 103 + idx
    sheet_proc.cell(row=current_row_proc, column=1, value=doc_num)
    sheet_proc.cell(row=current_row_proc, column=2, value=inv["Archivo"])
    sheet_proc.cell(row=current_row_proc, column=3, value=inv["Fecha"])
    sheet_proc.cell(row=current_row_proc, column=4, value=inv["Nº factura/documento"])
    sheet_proc.cell(row=current_row_proc, column=5, value=inv["Proveedor"])
    sheet_proc.cell(row=current_row_proc, column=6, value="ANA BELEN FEIJOO VIDUEIRA / 34975084L")
    sheet_proc.cell(row=current_row_proc, column=7, value=inv["Base imponible EUR"])
    sheet_proc.cell(row=current_row_proc, column=8, value=inv["Moneda original"])
    sheet_proc.cell(row=current_row_proc, column=9, value=inv["Estado revisión"])
    sheet_proc.cell(row=current_row_proc, column=10, value=inv["Observaciones"])
    current_row_proc += 1

wb.save("data/interim/ANA_FEIJOO_CALCULO_2T_2026.xlsx")
print("Appended received invoices.")
