import openpyxl
import pandas as pd
import os

os.makedirs("data/processed", exist_ok=True)
wb = openpyxl.load_workbook("data/interim/ANA_FEIJOO_CALCULO_2T_2026.xlsx")

# 1. Export "Libro Registro de Facturas Recibidas"
sheet_rec = wb["Facturas recibidas"]
rec_rows = []
for r in range(2, sheet_rec.max_row+1):
    row_val = [sheet_rec.cell(row=r, column=c).value for c in range(1, 25)]
    if any(x is not None for x in row_val):
        rec_rows.append(row_val)

cols_rec = ["Cliente", "NIF cliente", "Periodo", "Fecha", "Nº factura", "Proveedor", "NIF proveedor", "País", "Tipo operación", "Descripción", "Base", "IVA_rate", "IVA", "Total", "Moneda", "Imp_orig", "Deducibilidad", "Ret_rate", "Ret", "Clasif", "Estado", "Observaciones", "Doc_orig", "Incluye"]
df_rec = pd.DataFrame(rec_rows, columns=cols_rec)
df_rec_2t = df_rec[df_rec["Incluye"] == "Sí"].copy()
df_rec_2t = df_rec_2t.sort_values(by=["Fecha", "Nº factura"]).reset_index(drop=True)

official_rec_cols = {
    "Fecha Expedición/Recepción": df_rec_2t["Fecha"],
    "Número Factura/Documento": df_rec_2t["Nº factura"],
    "Proveedor (Emisor)": df_rec_2t["Proveedor"],
    "NIF/CIF Proveedor": df_rec_2t["NIF proveedor"],
    "País": df_rec_2t["País"],
    "Tipo de Operación": df_rec_2t["Tipo operación"],
    "Base Imponible (EUR)": df_rec_2t["Base"].astype(float).round(2),
    "Tipo de IVA (%)": (df_rec_2t["IVA_rate"].astype(float) * 100).round(2),
    "Cuota Soportada (EUR)": df_rec_2t["IVA"].astype(float).round(2),
    "Tipo de Retención (%)": (df_rec_2t["Ret_rate"].astype(float) * 100).round(2),
    "Retención IRPF (EUR)": df_rec_2t["Ret"].astype(float).round(2),
    "Total Factura (EUR)": df_rec_2t["Total"].astype(float).round(2),
    "Clasificación Fiscal": df_rec_2t["Clasif"],
    "Documento de Origen": df_rec_2t["Doc_orig"]
}
df_official_rec = pd.DataFrame(official_rec_cols)
df_official_rec.to_csv("data/processed/Libro_Registro_Facturas_Recibidas_2T_2026.csv", index=False, encoding="utf-8-sig", sep=";")
df_official_rec.to_excel("data/processed/Libro_Registro_Facturas_Recibidas_2T_2026.xlsx", index=False)

# 2. Export "Libro Registro de Facturas Emitidas"
sheet_em = wb["Facturas emitidas"]
em_rows = []
for r in range(4, sheet_em.max_row+1):
    row_val = [sheet_em.cell(row=r, column=c).value for c in range(1, 9)]
    if any(x is not None for x in row_val):
        em_rows.append(row_val)

cols_em = ["Fecha", "Factura", "Destinatario", "Base", "IVA", "Total", "Retencion", "Estado"]
df_em = pd.DataFrame(em_rows, columns=cols_em)
df_em = df_em.sort_values(by=["Fecha", "Factura"]).reset_index(drop=True)

official_em_cols = {
    "Fecha Expedición": df_em["Fecha"],
    "Número Factura/Referencia": df_em["Factura"],
    "Cliente/Destinatario": df_em["Destinatario"],
    "Base Imponible (EUR)": df_em["Base"].astype(float).round(2),
    "Cuota Repercutida (EUR)": df_em["IVA"].astype(float).round(2),
    "Retención IRPF (EUR)": df_em["Retencion"].astype(float).round(2),
    "Total Factura (EUR)": df_em["Total"].astype(float).round(2),
    "Estado de la Operación": df_em["Estado"]
}
df_official_em = pd.DataFrame(official_em_cols)
df_official_em.to_csv("data/processed/Libro_Registro_Facturas_Emitidas_2T_2026.csv", index=False, encoding="utf-8-sig", sep=";")
df_official_em.to_excel("data/processed/Libro_Registro_Facturas_Emitidas_2T_2026.xlsx", index=False)

print("Exported books.")
