import os
import openpyxl
import pytest

PROCESSED_DIR = "data/processed"

def test_files_exist():
    expected_files = [
        "Libro_Registro_Facturas_Emitidas_3T_2026.xlsx",
        "Libro_Registro_Facturas_Recibidas_3T_2026.xlsx",
        "Resumen_Fiscal_Modelo303_111_3T_2026.xlsx"
    ]
    for filename in expected_files:
        filepath = os.path.join(PROCESSED_DIR, filename)
        assert os.path.exists(filepath), f"File {filename} was not generated."

def test_facturas_emitidas_content():
    filepath = os.path.join(PROCESSED_DIR, "Libro_Registro_Facturas_Emitidas_3T_2026.xlsx")
    wb = openpyxl.load_workbook(filepath, data_only=False)
    sheet = wb.active

    # Check headers in row 4
    headers = [sheet.cell(row=4, column=col).value for col in range(1, 10)]
    assert "Nº Factura" in headers
    assert "Base Imponible (€)" in headers
    assert "Cuota IVA (€)" in headers

    # 14 issued invoices + 1 header + 3 title/empty + 1 total = row 19
    assert sheet.max_row == 19

    # Verify formula in row 5 for Cuota IVA
    assert sheet.cell(row=5, column=8).value == "=F5*G5"

    # Verify total row formula
    assert sheet.cell(row=19, column=6).value == "=SUM(F5:F18)"
    assert sheet.cell(row=19, column=8).value == "=SUM(H5:H18)"
    assert sheet.cell(row=19, column=9).value == "=SUM(I5:I18)"

def test_facturas_recibidas_content():
    filepath = os.path.join(PROCESSED_DIR, "Libro_Registro_Facturas_Recibidas_3T_2026.xlsx")
    wb = openpyxl.load_workbook(filepath, data_only=False)
    sheet = wb.active

    headers = [sheet.cell(row=4, column=col).value for col in range(1, 13)]
    assert "Nº Reg." in headers
    assert "NIF Proveedor" in headers
    assert "Retención IRPF (€)" in headers

    # 44 received items + 4 header/title rows + 1 total row = 49 rows
    assert sheet.max_row == 49

    # Check formula in row 5
    assert sheet.cell(row=5, column=9).value == "=G5*H5"
    assert sheet.cell(row=5, column=11).value == "=G5*J5"

    # Check total formulas
    assert sheet.cell(row=49, column=7).value == "=SUM(G5:G48)"
    assert sheet.cell(row=49, column=9).value == "=SUM(I5:I48)"
    assert sheet.cell(row=49, column=11).value == "=SUM(K5:K48)"
    assert sheet.cell(row=49, column=12).value == "=SUM(L5:L48)"

def test_resumen_fiscal_content():
    filepath = os.path.join(PROCESSED_DIR, "Resumen_Fiscal_Modelo303_111_3T_2026.xlsx")
    wb = openpyxl.load_workbook(filepath, data_only=False)

    assert "Modelo 303 (IVA 3T 2026)" in wb.sheetnames
    assert "Modelo 111 (IRPF 3T 2026)" in wb.sheetnames
    assert "Tributos Municipales 3T 2026" in wb.sheetnames

    ws_303 = wb["Modelo 303 (IVA 3T 2026)"]
    assert ws_303.cell(row=19, column=4).value == "=D17-D18"

    ws_111 = wb["Modelo 111 (IRPF 3T 2026)"]
    # 3 retainable invoices + 4 title/header + 1 total = 8 rows
    assert ws_111.max_row == 8
    assert ws_111.cell(row=8, column=8).value == "=SUM(H5:H7)"
