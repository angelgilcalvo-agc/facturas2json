import os
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build():
    # Load raw data
    source_path = '/tmp/file_attachments/2T 2026/ANA FEIJOO CALCULO 050726.xlsx'
    df = pd.read_excel(source_path, sheet_name='Facturas recibidas')

    # Clean completely empty rows
    df = df.dropna(how='all')

    # Check shape (should be 103 rows, 24 columns)
    print(f"Dataframe loaded: {df.shape}")

    # Sort by date for chronological accounting registry
    df['Fecha_dt'] = pd.to_datetime(df['Fecha'], errors='coerce')
    df = df.sort_values(by='Fecha_dt', na_position='last').drop(columns=['Fecha_dt'])

    # Reset index to have a neat sequence
    df = df.reset_index(drop=True)

    # Create output directory if it doesn't exist
    os.makedirs('data/processed', exist_ok=True)
    output_path = 'data/processed/Libro_Registro_Facturas_Recibidas_2T_2026.xlsx'

    # Build workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Facturas recibidas"

    # Define styles
    font_family = "Segoe UI"

    # Colors (Classic Navy Theme)
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")

    data_font = Font(name=font_family, size=10)
    totals_font = Font(name=font_family, size=10, bold=True)

    thin_side = Side(border_style="thin", color="D9D9D9")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    totals_top_side = Side(border_style="thin", color="000000")
    totals_bottom_side = Side(border_style="double", color="000000")
    totals_border = Border(top=totals_top_side, bottom=totals_bottom_side)

    align_left = Alignment(horizontal="left", vertical="center")
    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # Write headers
    headers = list(df.columns)
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = thin_border

    # Write data rows
    num_rows = len(df)
    for r_idx in range(num_rows):
        row_num = r_idx + 2
        row_data = df.iloc[r_idx]

        for c_idx, col_name in enumerate(headers, start=1):
            val = row_data[col_name]
            cell = ws.cell(row=row_num, column=c_idx)

            # Handle NaN / Null values
            if pd.isna(val):
                val_to_write = None
            else:
                val_to_write = val

            # Formats & alignments by column
            if col_name == 'Fecha':
                if val_to_write is not None:
                    # Parse as datetime if it is string
                    if isinstance(val_to_write, str):
                        cell.value = pd.to_datetime(val_to_write).date()
                    else:
                        cell.value = val_to_write
                else:
                    cell.value = None
                cell.number_format = 'dd/mm/yyyy'
                cell.alignment = align_center

            elif col_name in ['Base imponible EUR', 'IVA soportado EUR', 'Total EUR', 'Importe original', 'Retención EUR']:
                cell.value = float(val_to_write) if val_to_write is not None else None
                cell.number_format = '#,##0.00'
                cell.alignment = align_right

            elif col_name in ['% IVA', '% Retención']:
                cell.value = float(val_to_write) if val_to_write is not None else None
                cell.number_format = '0.0%'
                cell.alignment = align_right

            elif col_name == 'Deducibilidad':
                cell.value = float(val_to_write) if val_to_write is not None else None
                cell.number_format = '0%'
                cell.alignment = align_right

            elif col_name in ['NIF cliente', 'Nº factura/documento', 'NIF/CIF proveedor']:
                # Treat strictly as text
                cell.value = str(val_to_write) if val_to_write is not None else None
                cell.number_format = '@'
                cell.alignment = align_center

            else:
                # Text/general columns
                if val_to_write is not None:
                    cell.value = str(val_to_write)
                else:
                    cell.value = None
                cell.number_format = '@'
                cell.alignment = align_left

            cell.font = data_font
            cell.border = thin_border

    # Totals row
    totals_row_idx = num_rows + 2

    # Label "TOTAL GENERAL"
    total_label_cell = ws.cell(row=totals_row_idx, column=1, value="TOTAL GENERAL")
    total_label_cell.font = totals_font
    total_label_cell.alignment = align_left
    total_label_cell.border = totals_border

    # Fill borders for the entire totals row first
    for c_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=totals_row_idx, column=c_idx)
        cell.font = totals_font
        cell.border = totals_border
        if c_idx > 1 and cell.value is None:
            cell.value = None

    # Add formulas for numeric columns
    # Col K: Base imponible EUR (11)
    # Col M: IVA soportado EUR (13)
    # Col N: Total EUR (14)
    # Col P: Importe original (16)
    # Col S: Retención EUR (19)
    formula_cols = {
        11: 'K',
        13: 'M',
        14: 'N',
        16: 'P',
        19: 'S'
    }

    for col_idx, col_letter in formula_cols.items():
        cell = ws.cell(row=totals_row_idx, column=col_idx)
        cell.value = f"=SUM({col_letter}2:{col_letter}{totals_row_idx - 1})"
        cell.number_format = '#,##0.00'
        cell.alignment = align_right

    # Active AutoFilters on headers
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{totals_row_idx - 1}"

    # Freeze Header Pane (Freeze first row)
    ws.freeze_panes = 'A2'

    # Set Row Heights
    ws.row_dimensions[1].height = 28
    for r in range(2, totals_row_idx):
        ws.row_dimensions[r].height = 20
    ws.row_dimensions[totals_row_idx].height = 24

    # Auto-fit Column Widths (ignoring formulas for size calculation)
    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = len(headers[col_idx - 1])

        for r_idx in range(2, totals_row_idx):
            cell_val = ws.cell(row=r_idx, column=col_idx).value
            if cell_val is not None:
                max_len = max(max_len, len(str(cell_val)))

        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Save file
    wb.save(output_path)
    print(f"Workbook successfully saved to {output_path} with {num_rows} records.")

if __name__ == "__main__":
    build()
