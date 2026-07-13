import openpyxl
import os

def test_workbook_exists():
    assert os.path.exists("data/processed/informe_fiscal_gonzalo_2T_2026.xlsx")

def test_sheet_names():
    wb = openpyxl.load_workbook("data/processed/informe_fiscal_gonzalo_2T_2026.xlsx")
    expected_sheets = [
        "Datos del contribuyente",
        "Resumen ejecutivo",
        "Facturas emitidas",
        "Facturas recibidas",
        "Libro de ingresos",
        "Libro de gastos",
        "Bienes de inversión",
        "Retenciones",
        "IVA repercutido por tipos",
        "IVA soportado por tipos",
        "Cálculo modelo 303",
        "Cálculo modelo 130",
        "Resumen acumulado anual",
        "Operaciones con terceros",
        "Preparación modelo 347",
        "Preparación resúmenes anuales",
        "Alertas y revisión",
        "Documentos procesados",
        "Doc pendientes o rechazados"
    ]
    for s in expected_sheets:
        assert s in wb.sheetnames

def test_totals_and_formulas():
    wb = openpyxl.load_workbook("data/processed/informe_fiscal_gonzalo_2T_2026.xlsx")

    # Check Facturas emitidas totals formula
    ws = wb["Facturas emitidas"]
    assert ws["H16"].value == "=SUM(H2:H15)"
    assert ws["J16"].value == "=SUM(J2:J15)"
    assert ws["M16"].value == "=SUM(M2:M15)"
    assert ws["N16"].value == "=SUM(N2:N15)"

    # Check Libro de ingresos totals formula
    ws_ing = wb["Libro de ingresos"]
    assert ws_ing["G16"].value == "=SUM(G2:G15)"
    assert ws_ing["H16"].value == "=SUM(H2:H15)"
    assert ws_ing["I16"].value == "=SUM(I2:I15)"
    assert ws_ing["J16"].value == "=SUM(J2:J15)"

def test_casillas_formulas():
    wb = openpyxl.load_workbook("data/processed/informe_fiscal_gonzalo_2T_2026.xlsx")
    ws_303 = wb["Cálculo modelo 303"]
    # B8 is devengo total B5+B7
    assert ws_303["B8"].value == "=B5+B7"
    # B13 is total deducible B10+B12
    assert ws_303["B13"].value == "=B10+B12"
    # B17 is resultado B14-B15
    assert ws_303["B17"].value == "=B14-B15"

    ws_130 = wb["Cálculo modelo 130"]
    # B6 is rendimiento acumulado B4-B5
    assert ws_130["B6"].value == "=B4-B5"
    # B11 is resultado model 130 B7-B8-B9-B10
    assert ws_130["B11"].value == "=IF(B7-B8-B9-B10>0, B7-B8-B9-B10, 0)"
