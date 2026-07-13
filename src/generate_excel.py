import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

def create_tax_ledger_2026():
    wb = openpyxl.Workbook()
    # Remove default sheet
    default_sheet = wb.active
    wb.remove(default_sheet)

    # Styles
    font_title = Font(name="Segoe UI", size=16, bold=True, color="1F497D")
    font_subtitle = Font(name="Segoe UI", size=10, italic=True, color="595959")
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="Segoe UI", size=11, bold=False)
    font_total = Font(name="Segoe UI", size=11, bold=True, color="1F497D")

    fill_header = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    fill_total = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
    fill_alternate = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

    thin_border_side = Side(style='thin', color='E5E7EB')
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    total_border = Border(
        top=Side(style='thin', color='1F497D'),
        bottom=Side(style='double', color='1F497D'),
        left=thin_border_side, right=thin_border_side
    )

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    num_format_currency = '#,##0.00" €"'
    num_format_pct = '0.0%'
    num_format_date = 'YYYY-MM-DD'

    # We pre-format up to row 100 for easy manual additions
    max_data_row = 100
    total_row_idx = max_data_row + 1

    # 1. SHEET: FACTURAS EMITIDAS
    ws_emitidas = wb.create_sheet(title="Facturas Emitidas")
    ws_emitidas.views.sheetView[0].showGridLines = True

    # Title Block
    ws_emitidas["A1"] = "LIBRO REGISTRO DE FACTURAS EMITIDAS - EJERCICIO 2026"
    ws_emitidas["A1"].font = font_title
    ws_emitidas["A2"] = "Empresa: EXITER S.L. | NIF: B13658109"
    ws_emitidas["A2"].font = font_subtitle

    # Column Headers (Row 4)
    headers_emitidas = [
        "Trimestre", "Nº Factura", "Fecha Expedición", "NIF Cliente", "Cliente / Razón Social",
        "Base Imponible", "Tipo IVA", "Cuota IVA", "Total Factura"
    ]
    for col_idx, h in enumerate(headers_emitidas, 1):
        cell = ws_emitidas.cell(row=4, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border

    # Populate row 5 with extracted invoice
    invoice_date = datetime.strptime("2026-05-25", "%Y-%m-%d").date()
    data_emitidas_row5 = [
        "2T", "F0526016", invoice_date, "15.880.550-Q", "D.ª MARIASOL MÉNTRIDA LÓPEZ DE LAS HAZAS",
        1500.00, 0.21, 315.00, 1815.00
    ]
    for col_idx, val in enumerate(data_emitidas_row5, 1):
        cell = ws_emitidas.cell(row=5, column=col_idx, value=val)
        cell.font = font_data
        cell.border = thin_border

        # Alignment & formatting
        if col_idx in [1, 2, 4]:
            cell.alignment = align_center
        elif col_idx == 3:
            cell.alignment = align_center
            cell.number_format = num_format_date
        elif col_idx == 5:
            cell.alignment = align_left
        elif col_idx in [6, 8, 9]:
            cell.alignment = align_right
            cell.number_format = num_format_currency
        elif col_idx == 7:
            cell.alignment = align_right
            cell.number_format = num_format_pct

    # Pre-format rows 6 to 100 as blank rows
    for r_idx in range(6, max_data_row + 1):
        # alternate row background color for premium feel
        row_fill = fill_alternate if r_idx % 2 == 0 else None
        for col_idx in range(1, 10):
            cell = ws_emitidas.cell(row=r_idx, column=col_idx)
            cell.font = font_data
            cell.border = thin_border
            if row_fill:
                cell.fill = row_fill

            # Alignments & formats
            if col_idx in [1, 2, 4]:
                cell.alignment = align_center
            elif col_idx == 3:
                cell.alignment = align_center
                cell.number_format = num_format_date
            elif col_idx == 5:
                cell.alignment = align_left
            elif col_idx in [6, 8, 9]:
                cell.alignment = align_right
                cell.number_format = num_format_currency
            elif col_idx == 7:
                cell.alignment = align_right
                cell.number_format = num_format_pct

    # Total Row (Row 101)
    for col_idx in range(1, 5):
        cell = ws_emitidas.cell(row=total_row_idx, column=col_idx)
        cell.border = total_border
        cell.fill = fill_total

    cell_total_lbl = ws_emitidas.cell(row=total_row_idx, column=5, value="TOTAL")
    cell_total_lbl.font = font_total
    cell_total_lbl.alignment = align_right
    cell_total_lbl.border = total_border
    cell_total_lbl.fill = fill_total

    # Base Imponible sum
    cell_base_tot = ws_emitidas.cell(row=total_row_idx, column=6, value=f"=SUM(F5:F{max_data_row})")
    cell_base_tot.font = font_total
    cell_base_tot.alignment = align_right
    cell_base_tot.number_format = num_format_currency
    cell_base_tot.border = total_border
    cell_base_tot.fill = fill_total

    # Tipo IVA empty
    cell_tipo_tot = ws_emitidas.cell(row=total_row_idx, column=7)
    cell_tipo_tot.border = total_border
    cell_tipo_tot.fill = fill_total

    # Cuota IVA sum
    cell_cuota_tot = ws_emitidas.cell(row=total_row_idx, column=8, value=f"=SUM(H5:H{max_data_row})")
    cell_cuota_tot.font = font_total
    cell_cuota_tot.alignment = align_right
    cell_cuota_tot.number_format = num_format_currency
    cell_cuota_tot.border = total_border
    cell_cuota_tot.fill = fill_total

    # Total sum
    cell_total_tot = ws_emitidas.cell(row=total_row_idx, column=9, value=f"=SUM(I5:I{max_data_row})")
    cell_total_tot.font = font_total
    cell_total_tot.alignment = align_right
    cell_total_tot.number_format = num_format_currency
    cell_total_tot.border = total_border
    cell_total_tot.fill = fill_total


    # 2. SHEET: FACTURAS RECIBIDAS
    ws_recibidas = wb.create_sheet(title="Facturas Recibidas")
    ws_recibidas.views.sheetView[0].showGridLines = True

    # Title Block
    ws_recibidas["A1"] = "LIBRO REGISTRO DE FACTURAS RECIBIDAS - EJERCICIO 2026"
    ws_recibidas["A1"].font = font_title
    ws_recibidas["A2"] = "Empresa: EXITER S.L. | NIF: B13658109"
    ws_recibidas["A2"].font = font_subtitle

    # Column Headers (Row 4)
    headers_recibidas = [
        "Trimestre", "Nº Factura", "Fecha Expedición", "NIF Proveedor", "Proveedor / Razón Social",
        "Base Imponible", "Tipo IVA", "Cuota IVA", "Total Factura"
    ]
    for col_idx, h in enumerate(headers_recibidas, 1):
        cell = ws_recibidas.cell(row=4, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border

    # Pre-format rows 5 to 100 as blank rows
    for r_idx in range(5, max_data_row + 1):
        row_fill = fill_alternate if r_idx % 2 == 0 else None
        for col_idx in range(1, 10):
            cell = ws_recibidas.cell(row=r_idx, column=col_idx)
            cell.font = font_data
            cell.border = thin_border
            if row_fill:
                cell.fill = row_fill

            # Alignments & formats
            if col_idx in [1, 2, 4]:
                cell.alignment = align_center
            elif col_idx == 3:
                cell.alignment = align_center
                cell.number_format = num_format_date
            elif col_idx == 5:
                cell.alignment = align_left
            elif col_idx in [6, 8, 9]:
                cell.alignment = align_right
                cell.number_format = num_format_currency
            elif col_idx == 7:
                cell.alignment = align_right
                cell.number_format = num_format_pct

    # Total Row (Row 101)
    for col_idx in range(1, 5):
        cell = ws_recibidas.cell(row=total_row_idx, column=col_idx)
        cell.border = total_border
        cell.fill = fill_total

    cell_total_lbl_r = ws_recibidas.cell(row=total_row_idx, column=5, value="TOTAL")
    cell_total_lbl_r.font = font_total
    cell_total_lbl_r.alignment = align_right
    cell_total_lbl_r.border = total_border
    cell_total_lbl_r.fill = fill_total

    # Base Imponible sum
    cell_base_tot_r = ws_recibidas.cell(row=total_row_idx, column=6, value=f"=SUM(F5:F{max_data_row})")
    cell_base_tot_r.font = font_total
    cell_base_tot_r.alignment = align_right
    cell_base_tot_r.number_format = num_format_currency
    cell_base_tot_r.border = total_border
    cell_base_tot_r.fill = fill_total

    # Tipo IVA empty
    cell_tipo_tot_r = ws_recibidas.cell(row=total_row_idx, column=7)
    cell_tipo_tot_r.border = total_border
    cell_tipo_tot_r.fill = fill_total

    # Cuota IVA sum
    cell_cuota_tot_r = ws_recibidas.cell(row=total_row_idx, column=8, value=f"=SUM(H5:H{max_data_row})")
    cell_cuota_tot_r.font = font_total
    cell_cuota_tot_r.alignment = align_right
    cell_cuota_tot_r.number_format = num_format_currency
    cell_cuota_tot_r.border = total_border
    cell_cuota_tot_r.fill = fill_total

    # Total sum
    cell_total_tot_r = ws_recibidas.cell(row=total_row_idx, column=9, value=f"=SUM(I5:I{max_data_row})")
    cell_total_tot_r.font = font_total
    cell_total_tot_r.alignment = align_right
    cell_total_tot_r.number_format = num_format_currency
    cell_total_tot_r.border = total_border
    cell_total_tot_r.fill = fill_total


    # 3. SHEET: RESUMEN DE IMPUESTOS
    ws_resumen = wb.create_sheet(title="Resumen de Impuestos")
    ws_resumen.views.sheetView[0].showGridLines = True

    # Title Block
    ws_resumen["A1"] = "RESUMEN DE IMPUESTOS (IVA) - EJERCICIO 2026"
    ws_resumen["A1"].font = font_title
    ws_resumen["A2"] = "Empresa: EXITER S.L. | NIF: B13658109"
    ws_resumen["A2"].font = font_subtitle

    # --- TABLE 1: IVA REPERCUTIDO ---
    ws_resumen["A4"] = "1. IVA REPERCUTIDO (Ventas / Servicios Emitidos)"
    ws_resumen["A4"].font = Font(name="Segoe UI", size=12, bold=True, color="1F497D")

    headers_t1 = ["Periodo", "Base Imponible", "IVA Repercutido (Cuota)"]
    for col_idx, h in enumerate(headers_t1, 1):
        cell = ws_resumen.cell(row=5, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border

    periods = ["1T", "2T", "3T", "4T"]
    for idx, p in enumerate(periods, 6):
        # Period Label
        c1 = ws_resumen.cell(row=idx, column=1, value=p)
        c1.font = font_data
        c1.alignment = align_center
        c1.border = thin_border

        # Base Imponible formula
        c2 = ws_resumen.cell(row=idx, column=2, value=f"=SUMIFS('Facturas Emitidas'!F:F, 'Facturas Emitidas'!A:A, \"{p}\")")
        c2.font = font_data
        c2.alignment = align_right
        c2.number_format = num_format_currency
        c2.border = thin_border

        # Cuota IVA formula
        c3 = ws_resumen.cell(row=idx, column=3, value=f"=SUMIFS('Facturas Emitidas'!H:H, 'Facturas Emitidas'!A:A, \"{p}\")")
        c3.font = font_data
        c3.alignment = align_right
        c3.number_format = num_format_currency
        c3.border = thin_border

    # Total row for T1
    c1_tot = ws_resumen.cell(row=10, column=1, value="TOTAL ANUAL")
    c1_tot.font = font_total
    c1_tot.alignment = align_center
    c1_tot.border = total_border
    c1_tot.fill = fill_total

    c2_tot = ws_resumen.cell(row=10, column=2, value="=SUM(B6:B9)")
    c2_tot.font = font_total
    c2_tot.alignment = align_right
    c2_tot.number_format = num_format_currency
    c2_tot.border = total_border
    c2_tot.fill = fill_total

    c3_tot = ws_resumen.cell(row=10, column=3, value="=SUM(C6:C9)")
    c3_tot.font = font_total
    c3_tot.alignment = align_right
    c3_tot.number_format = num_format_currency
    c3_tot.border = total_border
    c3_tot.fill = fill_total


    # --- TABLE 2: IVA SOPORTADO ---
    ws_resumen["A12"] = "2. IVA SOPORTADO (Compras / Gastos Recibidos)"
    ws_resumen["A12"].font = Font(name="Segoe UI", size=12, bold=True, color="1F497D")

    headers_t2 = ["Periodo", "Base Imponible", "IVA Soportado (Cuota)"]
    for col_idx, h in enumerate(headers_t2, 1):
        cell = ws_resumen.cell(row=13, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border

    # Data rows for T2 (1T, 2T, 3T, 4T)
    for idx, p in enumerate(periods, 14):
        # Period Label
        c1 = ws_resumen.cell(row=idx, column=1, value=p)
        c1.font = font_data
        c1.alignment = align_center
        c1.border = thin_border

        # Base Imponible formula
        c2 = ws_resumen.cell(row=idx, column=2, value=f"=SUMIFS('Facturas Recibidas'!F:F, 'Facturas Recibidas'!A:A, \"{p}\")")
        c2.font = font_data
        c2.alignment = align_right
        c2.number_format = num_format_currency
        c2.border = thin_border

        # Cuota IVA formula
        c3 = ws_resumen.cell(row=idx, column=3, value=f"=SUMIFS('Facturas Recibidas'!H:H, 'Facturas Recibidas'!A:A, \"{p}\")")
        c3.font = font_data
        c3.alignment = align_right
        c3.number_format = num_format_currency
        c3.border = thin_border

    # Total row for T2
    c1_tot2 = ws_resumen.cell(row=18, column=1, value="TOTAL ANUAL")
    c1_tot2.font = font_total
    c1_tot2.alignment = align_center
    c1_tot2.border = total_border
    c1_tot2.fill = fill_total

    c2_tot2 = ws_resumen.cell(row=18, column=2, value="=SUM(B14:B17)")
    c2_tot2.font = font_total
    c2_tot2.alignment = align_right
    c2_tot2.number_format = num_format_currency
    c2_tot2.border = total_border
    c2_tot2.fill = fill_total

    c3_tot2 = ws_resumen.cell(row=18, column=3, value="=SUM(C14:C17)")
    c3_tot2.font = font_total
    c3_tot2.alignment = align_right
    c3_tot2.number_format = num_format_currency
    c3_tot2.border = total_border
    c3_tot2.fill = fill_total


    # --- TABLE 3: LIQUIDACIÓN DE IVA ---
    ws_resumen["A20"] = "3. LIQUIDACIÓN DE IVA (Resultado del Ejercicio)"
    ws_resumen["A20"].font = Font(name="Segoe UI", size=12, bold=True, color="1F497D")

    headers_t3 = ["Periodo", "IVA Repercutido", "IVA Soportado", "Resultado (A ingresar / A devolver)"]
    for col_idx, h in enumerate(headers_t3, 1):
        cell = ws_resumen.cell(row=21, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border

    for idx, p in enumerate(periods, 22):
        row_t1 = idx - 16  # row 22 corresponds to row 6 in T1
        row_t2 = idx - 8   # row 22 corresponds to row 14 in T2

        # Period Label
        c1 = ws_resumen.cell(row=idx, column=1, value=p)
        c1.font = font_data
        c1.alignment = align_center
        c1.border = thin_border

        # IVA Repercutido from Table 1 Cuota
        c2 = ws_resumen.cell(row=idx, column=2, value=f"=C{row_t1}")
        c2.font = font_data
        c2.alignment = align_right
        c2.number_format = num_format_currency
        c2.border = thin_border

        # IVA Soportado from Table 2 Cuota
        c3 = ws_resumen.cell(row=idx, column=3, value=f"=C{row_t2}")
        c3.font = font_data
        c3.alignment = align_right
        c3.number_format = num_format_currency
        c3.border = thin_border

        # Resultado (Repercutido - Soportado)
        c4 = ws_resumen.cell(row=idx, column=4, value=f"=B{idx}-C{idx}")
        c4.font = Font(name="Segoe UI", size=11, bold=True)
        c4.alignment = align_right
        c4.number_format = num_format_currency
        c4.border = thin_border

    # Total row for T3
    c1_tot3 = ws_resumen.cell(row=26, column=1, value="TOTAL ANUAL")
    c1_tot3.font = font_total
    c1_tot3.alignment = align_center
    c1_tot3.border = total_border
    c1_tot3.fill = fill_total

    c2_tot3 = ws_resumen.cell(row=26, column=2, value="=C10")
    c2_tot3.font = font_total
    c2_tot3.alignment = align_right
    c2_tot3.number_format = num_format_currency
    c2_tot3.border = total_border
    c2_tot3.fill = fill_total

    c3_tot3 = ws_resumen.cell(row=26, column=3, value="=C18")
    c3_tot3.font = font_total
    c3_tot3.alignment = align_right
    c3_tot3.number_format = num_format_currency
    c3_tot3.border = total_border
    c3_tot3.fill = fill_total

    c4_tot3 = ws_resumen.cell(row=26, column=4, value="=B26-C26")
    c4_tot3.font = font_total
    c4_tot3.alignment = align_right
    c4_tot3.number_format = num_format_currency
    c4_tot3.border = total_border
    c4_tot3.fill = fill_total


    # --- ADJUST COLUMN WIDTHS FOR ALL SHEETS ---
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = cell.value
                if val is not None:
                    val_str = str(val)
                    if not val_str.startswith("="):
                        if isinstance(val, datetime) or hasattr(val, 'strftime'):
                            val_str = val.strftime('%Y-%m-%d')
                        elif isinstance(val, float):
                            val_str = f"{val:,.2f} €"
                        max_len = max(max_len, len(val_str))
            sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

    ws_emitidas.column_dimensions['E'].width = 45
    ws_recibidas.column_dimensions['E'].width = 45
    ws_resumen.column_dimensions['A'].width = 45
    ws_resumen.column_dimensions['B'].width = 25
    ws_resumen.column_dimensions['C'].width = 25
    ws_resumen.column_dimensions['D'].width = 35

    # Ensure processed directory exists
    os.makedirs("data/processed", exist_ok=True)
    out_path = "data/processed/libro_registro_impuestos_2026.xlsx"
    wb.save(out_path)
    print(f"Workbook successfully saved to: {out_path}")

if __name__ == "__main__":
    create_tax_ledger_2026()
